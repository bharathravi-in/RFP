"""
Simple PPT Service - Clean, minimal PPT generation without templates.

This is a complete rewrite that:
- Never uses templates (always creates fresh presentations)
- Uses only direct shape creation
- Has conservative, safe dimensions
- Produces clean, professional slides
- Uses word wrap instead of truncation
"""
import io
import logging
from typing import Dict, List, Any, Optional

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

logger = logging.getLogger(__name__)


class SimplePPTService:
    """
    A simple, robust PPT generator that creates clean presentations
    without relying on templates or complex placeholder handling.
    """
    
    # Fixed slide dimensions (widescreen 16:9)
    SLIDE_WIDTH = Inches(13.333)
    SLIDE_HEIGHT = Inches(7.5)
    
    # Safe content area (with margins)
    MARGIN = 0.75  # inches
    SAFE_LEFT = MARGIN
    SAFE_TOP = 1.3  # Below header
    SAFE_WIDTH = 13.333 - (2 * MARGIN)  # ~11.8"
    SAFE_HEIGHT = 7.5 - 1.3 - MARGIN  # ~5.45"
    
    # Colors
    COLORS = {
        'primary': RGBColor(79, 70, 229),      # Indigo
        'secondary': RGBColor(99, 102, 241),   # Light indigo
        'accent': RGBColor(16, 185, 129),      # Green
        'text_dark': RGBColor(31, 41, 55),     # Dark gray
        'text_light': RGBColor(255, 255, 255), # White
        'bg_light': RGBColor(249, 250, 251),   # Light gray
        'red': RGBColor(220, 38, 38),          # Red
        'amber': RGBColor(245, 158, 11),       # Amber
        'purple': RGBColor(139, 92, 246),      # Purple
    }
    
    def __init__(self, branding: Optional[Dict] = None):
        """Initialize with optional branding colors."""
        if branding:
            if 'primary_color' in branding:
                self.COLORS['primary'] = self._hex_to_rgb(branding['primary_color'])
            if 'secondary_color' in branding:
                self.COLORS['secondary'] = self._hex_to_rgb(branding['secondary_color'])
    
    def _hex_to_rgb(self, hex_color: str) -> RGBColor:
        """Convert hex color to RGBColor."""
        hex_color = hex_color.lstrip('#')
        return RGBColor(
            int(hex_color[0:2], 16),
            int(hex_color[2:4], 16),
            int(hex_color[4:6], 16)
        )
    
    def _clean_bullets(self, bullets: List[str], max_count: int = 6) -> List[str]:
        """Clean bullet list - NO truncation, just clean markers."""
        if not bullets:
            return []
        result = []
        for b in bullets[:max_count]:
            if not b:
                continue
            # Remove any leading bullet markers
            clean = str(b).strip().lstrip('•-*→▪►◆').strip()
            clean = ' '.join(clean.split())  # Normalize spaces
            if clean:
                result.append(clean)
        return result
    
    def generate_pptx(
        self,
        slides_data: List[Dict[str, Any]],
        title: str = "Proposal",
        client_name: str = "Client",
        company_name: str = "Company"
    ) -> io.BytesIO:
        """
        Generate a PowerPoint presentation.
        Always creates a fresh presentation - no templates.
        """
        # Create fresh presentation
        prs = Presentation()
        prs.slide_width = self.SLIDE_WIDTH
        prs.slide_height = self.SLIDE_HEIGHT
        
        logger.info(f"Generating PPT with {len(slides_data)} slides (simple mode)")
        
        for slide_data in slides_data:
            slide_type = slide_data.get('slide_type', 'content')
            
            try:
                if slide_type == 'cover':
                    self._make_cover(prs, slide_data, title, client_name, company_name)
                elif slide_type == 'agenda':
                    self._make_agenda(prs, slide_data)
                elif slide_type == 'architecture':
                    self._make_architecture(prs, slide_data)
                elif slide_type == 'timeline':
                    self._make_timeline(prs, slide_data)
                elif slide_type == 'team':
                    self._make_team(prs, slide_data)
                elif slide_type == 'pricing':
                    self._make_pricing(prs, slide_data)
                elif slide_type == 'two_column':
                    self._make_two_column(prs, slide_data)
                elif slide_type == 'problem':
                    self._make_content(prs, slide_data, header_color=self.COLORS['red'])
                elif slide_type == 'solution':
                    self._make_content(prs, slide_data, header_color=self.COLORS['accent'])
                elif slide_type == 'risk':
                    self._make_content(prs, slide_data, header_color=self.COLORS['amber'])
                elif slide_type == 'roi':
                    self._make_content(prs, slide_data, header_color=self.COLORS['purple'])
                elif slide_type == 'closing':
                    self._make_closing(prs, slide_data, company_name)
                else:
                    self._make_content(prs, slide_data)
            except Exception as e:
                logger.error(f"Error creating {slide_type} slide: {e}")
                import traceback
                traceback.print_exc()
                # Create fallback content slide
                self._make_content(prs, slide_data)
        
        # Save to buffer
        buffer = io.BytesIO()
        prs.save(buffer)
        buffer.seek(0)
        
        logger.info(f"Generated PPT successfully with {len(prs.slides)} slides")
        return buffer
    
    def _add_header(self, slide, title: str, color: RGBColor = None):
        """Add a standard header bar with title."""
        color = color or self.COLORS['primary']
        
        # Header bar
        header = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            Inches(0), Inches(0),
            self.SLIDE_WIDTH, Inches(1.1)
        )
        header.fill.solid()
        header.fill.fore_color.rgb = color
        header.line.fill.background()
        
        # Title text with word wrap
        title_box = slide.shapes.add_textbox(
            Inches(0.75), Inches(0.25),
            Inches(11.5), Inches(0.7)
        )
        tf = title_box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = title or 'Content'
        p.font.size = Pt(26)
        p.font.bold = True
        p.font.color.rgb = self.COLORS['text_light']
    
    def _add_bullets(self, slide, bullets: List[str], top: float = 1.4, font_size: int = None):
        """Add bullet points to slide with word wrap."""
        if not bullets:
            return
        
        content_box = slide.shapes.add_textbox(
            Inches(self.SAFE_LEFT), Inches(top),
            Inches(self.SAFE_WIDTH), Inches(5.5)
        )
        tf = content_box.text_frame
        tf.word_wrap = True
        
        # Calculate font size based on content length
        if font_size is None:
            total_len = sum(len(b) for b in bullets)
            if total_len > 500 or len(bullets) > 6:
                font_size = 14
            elif total_len > 350 or len(bullets) > 5:
                font_size = 16
            elif total_len > 200 or len(bullets) > 4:
                font_size = 18
            else:
                font_size = 20
        
        for i, bullet in enumerate(bullets):
            if i == 0:
                p = tf.paragraphs[0]
            else:
                p = tf.add_paragraph()
            
            p.text = f"• {bullet}"
            p.font.size = Pt(font_size)
            p.font.color.rgb = self.COLORS['text_dark']
            p.space_before = Pt(10)
            p.space_after = Pt(4)
    
    # ===== SLIDE TYPES =====
    
    def _make_cover(self, prs, data, title, client_name, company_name):
        """Create cover slide."""
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        
        # Full background
        bg = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            Inches(0), Inches(0),
            self.SLIDE_WIDTH, self.SLIDE_HEIGHT
        )
        bg.fill.solid()
        bg.fill.fore_color.rgb = self.COLORS['primary']
        bg.line.fill.background()
        
        # Title with word wrap
        title_box = slide.shapes.add_textbox(
            Inches(1), Inches(2.2),
            Inches(11), Inches(2)
        )
        tf = title_box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = data.get('title', title)
        p.font.size = Pt(36)
        p.font.bold = True
        p.font.color.rgb = self.COLORS['text_light']
        p.alignment = PP_ALIGN.CENTER
        
        # Subtitle with word wrap
        sub_box = slide.shapes.add_textbox(
            Inches(1), Inches(4.5),
            Inches(11), Inches(1)
        )
        tf = sub_box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = data.get('subtitle', f"Proposal for {client_name}")
        p.font.size = Pt(18)
        p.font.color.rgb = self.COLORS['text_light']
        p.alignment = PP_ALIGN.CENTER
        
        # Company name at bottom
        company_box = slide.shapes.add_textbox(
            Inches(1), Inches(6.5),
            Inches(11), Inches(0.5)
        )
        tf = company_box.text_frame
        p = tf.paragraphs[0]
        p.text = company_name
        p.font.size = Pt(14)
        p.font.color.rgb = self.COLORS['text_light']
        p.alignment = PP_ALIGN.CENTER
    
    def _make_agenda(self, prs, data):
        """Create agenda slide."""
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        self._add_header(slide, data.get('title', 'Agenda'))
        
        bullets = self._clean_bullets(data.get('bullets', []), max_count=8)
        self._add_bullets(slide, bullets)
    
    def _make_content(self, prs, data, header_color=None):
        """Create standard content slide."""
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        self._add_header(slide, data.get('title', 'Content'), header_color)
        
        bullets = self._clean_bullets(data.get('bullets', []), max_count=7)
        self._add_bullets(slide, bullets)
    
    def _make_architecture(self, prs, data):
        """Create architecture slide with diagram or layered boxes."""
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        self._add_header(slide, data.get('title', 'Solution Architecture'))
        
        # Try to render mermaid diagram if available
        mermaid_code = data.get('mermaid_code')
        diagram_rendered = False
        
        if mermaid_code:
            try:
                from .mermaid_service import render_mermaid_to_bytes_io
                diagram_buffer = render_mermaid_to_bytes_io(mermaid_code)
                if diagram_buffer:
                    # Add the diagram image
                    slide.shapes.add_picture(
                        diagram_buffer,
                        Inches(1), Inches(1.5),
                        width=Inches(11)
                    )
                    diagram_rendered = True
                    logger.info("Architecture diagram rendered successfully")
            except Exception as e:
                logger.warning(f"Could not render mermaid diagram: {e}")
        
        if not diagram_rendered:
            # Fallback: Draw layered architecture boxes
            bullets = self._clean_bullets(data.get('bullets', []), max_count=4)
            
            if not bullets or len(bullets) < 2:
                bullets = [
                    'Presentation Layer: Web UI, Mobile App, Admin Portal',
                    'Application Layer: Core Services, Business Logic, Workflows',
                    'Integration Layer: REST APIs, External Systems, Authentication',
                    'Data Layer: Database, Document Storage, Analytics Engine'
                ]
            
            layer_colors = [
                RGBColor(99, 102, 241),
                RGBColor(79, 70, 229),
                RGBColor(67, 56, 202),
                RGBColor(55, 48, 163),
            ]
            
            start_y = 1.5
            layer_height = 1.1
            gap = 0.12
            
            for i, layer_text in enumerate(bullets[:4]):
                y = start_y + i * (layer_height + gap)
                
                # Layer box
                box = slide.shapes.add_shape(
                    MSO_SHAPE.ROUNDED_RECTANGLE,
                    Inches(1), Inches(y),
                    Inches(11), Inches(layer_height)
                )
                box.fill.solid()
                box.fill.fore_color.rgb = layer_colors[i % 4]
                box.line.fill.background()
                
                # Layer text with word wrap
                text_box = slide.shapes.add_textbox(
                    Inches(1.2), Inches(y + 0.25),
                    Inches(10.6), Inches(0.7)
                )
                tf = text_box.text_frame
                tf.word_wrap = True
                p = tf.paragraphs[0]
                p.text = layer_text
                p.font.size = Pt(15)
                p.font.bold = True
                p.font.color.rgb = self.COLORS['text_light']
    
    def _make_timeline(self, prs, data):
        """Create timeline slide with horizontal phases."""
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        self._add_header(slide, data.get('title', 'Project Timeline'))
        
        phases = self._clean_bullets(data.get('bullets', []), max_count=5)
        
        if not phases:
            phases = ['Initiate & Plan', 'Design', 'Develop', 'Test', 'Deploy']
        
        num = min(len(phases), 5)
        phases = phases[:num]
        
        # Calculate positions
        total_width = 11.0
        start_x = 1.0
        phase_width = total_width / num
        
        # Connector line
        line_y = 3.0
        connector = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            Inches(start_x + 0.3), Inches(line_y),
            Inches(total_width - 0.6), Inches(0.06)
        )
        connector.fill.solid()
        connector.fill.fore_color.rgb = RGBColor(209, 213, 219)
        connector.line.fill.background()
        
        colors = [
            self.COLORS['primary'],
            self.COLORS['secondary'],
            self.COLORS['accent'],
            self.COLORS['purple'],
            self.COLORS['amber'],
        ]
        
        for i, phase in enumerate(phases):
            x = start_x + i * phase_width
            
            # Circle
            circle = slide.shapes.add_shape(
                MSO_SHAPE.OVAL,
                Inches(x + phase_width/2 - 0.25), Inches(line_y - 0.22),
                Inches(0.5), Inches(0.5)
            )
            circle.fill.solid()
            circle.fill.fore_color.rgb = colors[i % 5]
            circle.line.color.rgb = self.COLORS['text_light']
            circle.line.width = Pt(2)
            
            # Number
            num_box = slide.shapes.add_textbox(
                Inches(x + phase_width/2 - 0.25), Inches(line_y - 0.17),
                Inches(0.5), Inches(0.4)
            )
            tf = num_box.text_frame
            p = tf.paragraphs[0]
            p.text = str(i + 1)
            p.font.size = Pt(14)
            p.font.bold = True
            p.font.color.rgb = self.COLORS['text_light']
            p.alignment = PP_ALIGN.CENTER
            
            # Label with word wrap
            label_box = slide.shapes.add_textbox(
                Inches(x + 0.1), Inches(line_y + 0.5),
                Inches(phase_width - 0.2), Inches(1.5)
            )
            tf = label_box.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.text = phase
            p.font.size = Pt(11)
            p.font.bold = True
            p.font.color.rgb = self.COLORS['text_dark']
            p.alignment = PP_ALIGN.CENTER
    
    def _make_team(self, prs, data):
        """Create team slide with role cards."""
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        self._add_header(slide, data.get('title', 'Project Team'))
        
        members = self._clean_bullets(data.get('bullets', []), max_count=6)
        
        if not members:
            members = ['Project Manager', 'Technical Lead', 'Solution Architect', 'Developer']
        
        num = min(len(members), 6)
        members = members[:num]
        
        # Grid layout
        cols = min(3, num)
        
        # Card dimensions
        card_w = 3.4
        card_h = 2.0
        gap_x = 0.25
        gap_y = 0.2
        
        total_w = cols * card_w + (cols - 1) * gap_x
        start_x = (13.333 - total_w) / 2
        start_y = 1.4
        
        for i, member in enumerate(members):
            col = i % cols
            row = i // cols
            
            x = start_x + col * (card_w + gap_x)
            y = start_y + row * (card_h + gap_y)
            
            # Card
            card = slide.shapes.add_shape(
                MSO_SHAPE.ROUNDED_RECTANGLE,
                Inches(x), Inches(y),
                Inches(card_w), Inches(card_h)
            )
            card.fill.solid()
            card.fill.fore_color.rgb = self.COLORS['bg_light']
            card.line.color.rgb = RGBColor(209, 213, 219)
            
            # Number circle
            circle = slide.shapes.add_shape(
                MSO_SHAPE.OVAL,
                Inches(x + card_w/2 - 0.22), Inches(y + 0.2),
                Inches(0.44), Inches(0.44)
            )
            circle.fill.solid()
            circle.fill.fore_color.rgb = self.COLORS['secondary']
            circle.line.fill.background()
            
            # Number
            num_box = slide.shapes.add_textbox(
                Inches(x + card_w/2 - 0.22), Inches(y + 0.26),
                Inches(0.44), Inches(0.36)
            )
            tf = num_box.text_frame
            p = tf.paragraphs[0]
            p.text = str(i + 1)
            p.font.size = Pt(12)
            p.font.bold = True
            p.font.color.rgb = self.COLORS['text_light']
            p.alignment = PP_ALIGN.CENTER
            
            # Role text with word wrap
            text_box = slide.shapes.add_textbox(
                Inches(x + 0.15), Inches(y + 0.75),
                Inches(card_w - 0.3), Inches(1.1)
            )
            tf = text_box.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.text = member
            p.font.size = Pt(11)
            p.font.bold = True
            p.font.color.rgb = self.COLORS['text_dark']
            p.alignment = PP_ALIGN.CENTER
    
    def _make_pricing(self, prs, data):
        """Create pricing slide."""
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        self._add_header(slide, data.get('title', 'Investment Summary'), self.COLORS['accent'])
        
        bullets = self._clean_bullets(data.get('bullets', []), max_count=6)
        self._add_bullets(slide, bullets, font_size=18)
    
    def _make_two_column(self, prs, data):
        """Create two-column slide."""
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        self._add_header(slide, data.get('title', 'Content'))
        
        bullets = data.get('bullets', [])
        
        # Split bullets into two columns
        mid = len(bullets) // 2 if len(bullets) > 1 else len(bullets)
        left_bullets = self._clean_bullets(bullets[:mid], max_count=4)
        right_bullets = self._clean_bullets(bullets[mid:], max_count=4)
        
        # Left column
        if left_bullets:
            left_box = slide.shapes.add_textbox(
                Inches(0.75), Inches(1.4),
                Inches(5.5), Inches(5)
            )
            tf = left_box.text_frame
            tf.word_wrap = True
            
            for i, bullet in enumerate(left_bullets):
                if i == 0:
                    p = tf.paragraphs[0]
                else:
                    p = tf.add_paragraph()
                p.text = f"• {bullet}"
                p.font.size = Pt(16)
                p.font.color.rgb = self.COLORS['text_dark']
                p.space_before = Pt(10)
        
        # Right column
        if right_bullets:
            right_box = slide.shapes.add_textbox(
                Inches(6.75), Inches(1.4),
                Inches(5.5), Inches(5)
            )
            tf = right_box.text_frame
            tf.word_wrap = True
            
            for i, bullet in enumerate(right_bullets):
                if i == 0:
                    p = tf.paragraphs[0]
                else:
                    p = tf.add_paragraph()
                p.text = f"• {bullet}"
                p.font.size = Pt(16)
                p.font.color.rgb = self.COLORS['text_dark']
                p.space_before = Pt(10)
    
    def _make_closing(self, prs, data, company_name):
        """Create closing slide."""
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        
        # Full background
        bg = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            Inches(0), Inches(0),
            self.SLIDE_WIDTH, self.SLIDE_HEIGHT
        )
        bg.fill.solid()
        bg.fill.fore_color.rgb = self.COLORS['primary']
        bg.line.fill.background()
        
        # Thank you
        title_box = slide.shapes.add_textbox(
            Inches(1), Inches(2.5),
            Inches(11), Inches(1)
        )
        tf = title_box.text_frame
        p = tf.paragraphs[0]
        p.text = data.get('title', 'Thank You')
        p.font.size = Pt(44)
        p.font.bold = True
        p.font.color.rgb = self.COLORS['text_light']
        p.alignment = PP_ALIGN.CENTER
        
        # Subtitle
        sub_box = slide.shapes.add_textbox(
            Inches(1), Inches(4),
            Inches(11), Inches(0.6)
        )
        tf = sub_box.text_frame
        p = tf.paragraphs[0]
        p.text = "Questions & Discussion"
        p.font.size = Pt(22)
        p.font.color.rgb = self.COLORS['text_light']
        p.alignment = PP_ALIGN.CENTER
        
        # Company
        company_box = slide.shapes.add_textbox(
            Inches(1), Inches(6),
            Inches(11), Inches(0.5)
        )
        tf = company_box.text_frame
        p = tf.paragraphs[0]
        p.text = company_name
        p.font.size = Pt(14)
        p.font.color.rgb = self.COLORS['text_light']
        p.alignment = PP_ALIGN.CENTER
